from datetime import datetime, date, time
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from core.models import OrganizationUnit

from .models import (
    Assessment,
    AssessmentIndicator,
    AssessmentPeriod,
    AssessmentResponse,
    AssessmentSection,
)

# ================================================================
# EXPORT LIBRARIES
# ================================================================

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    SimpleDocTemplate,
    PageBreak,
)


# ================================================================
# PROFILE
# ================================================================

def get_active_profile(request):
    """
    Return the active EHIRS profile for the logged-in user.
    """

    try:
        profile = request.user.profile
    except Exception:
        return None

    if not profile:
        return None

    if not profile.active:
        return None

    return profile


# ================================================================
# ACCESSIBLE FACILITIES
# ================================================================

def get_accessible_facilities(profile):
    """
    Return only active facilities accessible to the current user.

    Facility user:
        Own facility only.

    Sub-city:
        Facilities belonging to own sub-city.

    Region:
        Facilities in own region, including facilities
        belonging to its sub-cities.

    Ministry:
        All facilities allowed by UserProfile.
    """

    return (
        profile.accessible_facilities()
        .select_related(
            "organization",
            "organization__parent",
            "organization__parent__parent",
        )
        .filter(active=True)
        .order_by(
            "organization__name",
            "name",
        )
    )


# ================================================================
# ACCESS CHECK
# ================================================================

def user_can_access_assessment(profile, assessment):
    """
    Verify that the user can access the assessment's facility.
    """

    if not profile:
        return False

    if not profile.active:
        return False

    if not assessment:
        return False

    return profile.can_access_facility(
        assessment.facility
    )


# ================================================================
# EDIT CHECK
# ================================================================

def user_can_edit_assessment(profile, assessment):
    """
    Only the facility user assigned to the facility can edit.

    Editable statuses:
        DRAFT
        RETURNED
    """

    if not profile:
        return False

    if not profile.active:
        return False

    if not profile.is_facility_user():
        return False

    if not profile.facility_id:
        return False

    if assessment.facility_id != profile.facility_id:
        return False

    return assessment.status in [
        Assessment.Status.DRAFT,
        Assessment.Status.RETURNED,
    ]


# ================================================================
# REVIEW PERMISSION
# ================================================================

def user_can_review_assessment(profile, assessment):
    """
    Ministry, Region and Sub-city users can review.

    Facility users cannot review.
    """

    if not profile:
        return False

    if not profile.active:
        return False

    if profile.is_facility_user():
        return False

    return profile.can_access_facility(
        assessment.facility
    )


# ================================================================
# APPROVAL PERMISSION
# ================================================================

def user_can_approve_assessment(profile, assessment):
    """
    Ministry and Region users can approve.

    Sub-city users cannot approve.
    Facility users cannot approve.
    """

    if not profile:
        return False

    if not profile.active:
        return False

    if profile.is_facility_user():
        return False

    if profile.is_subcity_admin():
        return False

    if not (
        profile.is_ministry_admin()
        or profile.is_region_admin()
    ):
        return False

    return profile.can_access_facility(
        assessment.facility
    )


# ================================================================
# ORGANIZATION HIERARCHY
# ================================================================

def build_hierarchy(organization):
    """
    Return organization hierarchy from top to bottom.
    """

    hierarchy = []

    current = organization

    while current:
        hierarchy.insert(
            0,
            current,
        )

        current = current.parent

    return hierarchy


def get_region_for_facility(facility):
    """
    Find the region belonging to a facility.
    """

    if not facility or not facility.organization:
        return None

    hierarchy = build_hierarchy(
        facility.organization
    )

    for organization in hierarchy:

        if (
            organization.unit_type
            == OrganizationUnit.UnitType.REGION
        ):
            return organization

    return None


def get_subcity_for_facility(facility):
    """
    Find the sub-city belonging to a facility.
    """

    if not facility or not facility.organization:
        return None

    hierarchy = build_hierarchy(
        facility.organization
    )

    for organization in hierarchy:

        if (
            organization.unit_type
            == OrganizationUnit.UnitType.SUBCITY
        ):
            return organization

    return None


# ================================================================
# ACCESSIBLE REGIONS
# ================================================================

def get_accessible_regions(profile):
    """
    Regions visible to the logged-in user.
    """

    if profile.is_ministry_admin():

        return (
            OrganizationUnit.objects
            .filter(
                active=True,
                unit_type=OrganizationUnit.UnitType.REGION,
            )
            .order_by("name")
        )

    if profile.is_region_admin():

        if (
            profile.organization
            and profile.organization.unit_type
            == OrganizationUnit.UnitType.REGION
        ):

            return (
                OrganizationUnit.objects
                .filter(
                    id=profile.organization_id,
                    active=True,
                )
            )

    return OrganizationUnit.objects.none()


# ================================================================
# ACCESSIBLE SUB-CITIES
# ================================================================

def get_accessible_subcities(
    profile,
    selected_region=None,
):
    """
    Return sub-cities visible to the user.
    """

    queryset = (
        OrganizationUnit.objects
        .filter(
            active=True,
            unit_type=OrganizationUnit.UnitType.SUBCITY,
        )
        .select_related("parent")
        .order_by("name")
    )

    # ------------------------------------------------------------
    # MINISTRY
    # ------------------------------------------------------------

    if profile.is_ministry_admin():

        if selected_region:

            return queryset.filter(
                parent=selected_region
            )

        return queryset

    # ------------------------------------------------------------
    # REGION
    # ------------------------------------------------------------

    if profile.is_region_admin():

        if not profile.organization:

            return queryset.none()

        return queryset.filter(
            parent=profile.organization
        )

    # ------------------------------------------------------------
    # SUB-CITY
    # ------------------------------------------------------------

    if profile.is_subcity_admin():

        if (
            profile.organization
            and profile.organization.unit_type
            == OrganizationUnit.UnitType.SUBCITY
        ):

            return queryset.filter(
                id=profile.organization_id
            )

    return queryset.none()


# ================================================================
# PERIODS
# ================================================================

def get_active_periods():
    """
    Return active assessment periods.
    """

    return (
        AssessmentPeriod.objects
        .filter(active=True)
        .order_by(
            "-year",
            "start_month",
            "period_number",
        )
    )


def get_years():
    """
    Return available assessment years.
    """

    periods = get_active_periods()

    years = list(
        periods
        .values_list(
            "year",
            flat=True,
        )
        .distinct()
    )

    return sorted(
        years,
        reverse=True,
    )


# ================================================================
# START ASSESSMENT
# ================================================================

@login_required
def assessment_start(request):
    """
    Start an assessment.

    Facility user:
        Own facility is automatically selected.

    Higher-level users:
        Can select an accessible facility.

    Year and reporting period are selected before
    opening the checklist.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    facilities = get_accessible_facilities(
        profile
    )

    periods = get_active_periods()

    years = get_years()

    facility_user = profile.is_facility_user()

    assigned_facility = None

    # ============================================================
    # FACILITY USER
    # ============================================================

    if facility_user:

        if not profile.facility_id:

            messages.error(
                request,
                "No facility is assigned to your account.",
            )

            return redirect(
                "assessments:dashboard"
            )

        assigned_facility = profile.facility

        if not assigned_facility.active:

            messages.error(
                request,
                "Your assigned facility is inactive.",
            )

            return redirect(
                "assessments:dashboard"
            )

    # ============================================================
    # POST
    # ============================================================

    if request.method == "POST":

        # --------------------------------------------------------
        # FACILITY
        # --------------------------------------------------------

        if facility_user:

            facility = assigned_facility

        else:

            facility_id = request.POST.get(
                "facility"
            )

            if not facility_id:

                messages.error(
                    request,
                    "Please select a facility.",
                )

                return render(
                    request,
                    "assessments/assessment_start.html",
                    {
                        "profile": profile,
                        "facilities": facilities,
                        "periods": periods,
                        "years": years,
                        "facility_user": facility_user,
                        "assigned_facility": assigned_facility,
                    },
                )

            facility = get_object_or_404(
                facilities,
                id=facility_id,
                active=True,
            )

        # --------------------------------------------------------
        # YEAR
        # --------------------------------------------------------

        year_value = request.POST.get(
            "year"
        )

        try:

            selected_year = int(
                year_value
            )

        except (
            TypeError,
            ValueError,
        ):

            messages.error(
                request,
                "Please select a valid year.",
            )

            return redirect(
                "assessments:assessment_start"
            )

        # --------------------------------------------------------
        # PERIOD
        # --------------------------------------------------------

        period_id = request.POST.get(
            "period"
        )

        if not period_id:

            messages.error(
                request,
                "Please select a reporting period.",
            )

            return redirect(
                "assessments:assessment_start"
            )

        period = get_object_or_404(
            periods,
            id=period_id,
            year=selected_year,
        )

        # --------------------------------------------------------
        # CREATE / OPEN
        # --------------------------------------------------------

        assessment, created = (
            Assessment.objects.get_or_create(
                facility=facility,
                period=period,
                defaults={
                    "created_by": request.user,
                },
            )
        )

        # --------------------------------------------------------
        # EXISTING ASSESSMENT
        # --------------------------------------------------------

        if not created:

            if assessment.status not in [
                Assessment.Status.DRAFT,
                Assessment.Status.RETURNED,
            ]:

                messages.info(
                    request,
                    (
                        "This assessment already exists and is "
                        f"{assessment.get_status_display()}. "
                        "It will be opened in view-only mode."
                    ),
                )

        return redirect(
            "assessments:assessment_detail",
            assessment_id=assessment.pk,
        )

    # ============================================================
    # DISPLAY
    # ============================================================

    return render(
        request,
        "assessments/assessment_start.html",
        {
            "profile": profile,
            "facilities": facilities,
            "periods": periods,
            "years": years,
            "facility_user": facility_user,
            "assigned_facility": assigned_facility,
        },
    )


# ================================================================
# ASSESSMENT DETAIL
# ================================================================

@login_required
def assessment_detail(
    request,
    assessment_id,
):
    """
    Display an assessment.

    Facility users can edit their own DRAFT/RETURNED assessment.

    Higher-level users can view assessments within their
    permitted organizational scope.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    assessment = get_object_or_404(
        Assessment.objects.select_related(
            "facility",
            "facility__organization",
            "period",
            "created_by",
        ),
        pk=assessment_id,
    )

    # ============================================================
    # ACCESS CONTROL
    # ============================================================

    if not user_can_access_assessment(
        profile,
        assessment,
    ):

        messages.error(
            request,
            "You do not have permission to access this assessment.",
        )

        return redirect(
            "assessments:dashboard"
        )

    # ============================================================
    # EDIT PERMISSION
    # ============================================================

    can_edit = user_can_edit_assessment(
        profile,
        assessment,
    )

    # ============================================================
    # REVIEW / APPROVAL
    # ============================================================

    can_review = user_can_review_assessment(
        profile,
        assessment
    )

    can_approve = user_can_approve_assessment(
        profile,
        assessment
    )

    # ============================================================
    # STRUCTURE
    # ============================================================

    sections = (
        AssessmentSection.objects
        .filter(active=True)
        .prefetch_related(
            "thematic_areas__indicators__options"
        )
        .order_by(
            "order",
            "code",
        )
    )

    # ============================================================
    # EXISTING RESPONSES
    # ============================================================

    existing_responses = {
        response.indicator_id: response
        for response in (
            assessment.responses
            .select_related(
                "selected_option",
                "indicator",
            )
        )
    }

    # ============================================================
    # POST
    # ============================================================

    if request.method == "POST":

        # --------------------------------------------------------
        # SECURITY
        # --------------------------------------------------------

        if not can_edit:

            messages.error(
                request,
                "This assessment cannot be edited.",
            )

            return redirect(
                "assessments:assessment_detail",
                assessment_id=assessment.pk,
            )

        action = (
            request.POST.get(
                "action",
                "save",
            )
            .strip()
            .lower()
        )

        if action not in {
            "save",
            "submit",
        }:

            messages.error(
                request,
                "Invalid assessment action.",
            )

            return redirect(
                "assessments:assessment_detail",
                assessment_id=assessment.pk,
            )

        # ========================================================
        # SAVE RESPONSES
        # ========================================================

        try:

            with transaction.atomic():

                for section in sections:

                    for thematic_area in (
                        section.thematic_areas
                        .filter(active=True)
                    ):

                        indicators = (
                            thematic_area.indicators
                            .filter(active=True)
                            .prefetch_related("options")
                            .order_by(
                                "order",
                                "code",
                            )
                        )

                        for indicator in indicators:

                            prefix = (
                                f"indicator_{indicator.id}"
                            )

                            # ------------------------------------
                            # POST VALUES
                            # ------------------------------------

                            option_id = request.POST.get(
                                f"{prefix}_option"
                            )

                            numeric_value = request.POST.get(
                                f"{prefix}_numeric"
                            )

                            numerator = request.POST.get(
                                f"{prefix}_numerator"
                            )

                            denominator = request.POST.get(
                                f"{prefix}_denominator"
                            )

                            conducted = request.POST.get(
                                f"{prefix}_conducted"
                            )

                            notes = request.POST.get(
                                f"{prefix}_notes",
                                "",
                            ).strip()

                            # ------------------------------------
                            # RESPONSE
                            # ------------------------------------

                            response, _ = (
                                AssessmentResponse.objects
                                .get_or_create(
                                    assessment=assessment,
                                    indicator=indicator,
                                )
                            )

                            # ------------------------------------
                            # OPTION
                            # ------------------------------------

                            selected_option = None

                            if option_id:

                                selected_option = (
                                    indicator.options
                                    .filter(
                                        id=option_id,
                                        active=True,
                                    )
                                    .first()
                                )

                            response.selected_option = (
                                selected_option
                            )

                            # ------------------------------------
                            # DECIMAL HELPER
                            # ------------------------------------

                            def decimal_or_none(value):

                                if value in [
                                    None,
                                    "",
                                ]:
                                    return None

                                try:

                                    return Decimal(
                                        value
                                    )

                                except (
                                    InvalidOperation,
                                    TypeError,
                                    ValueError,
                                ):

                                    return None

                            # ------------------------------------
                            # NUMERIC
                            # ------------------------------------

                            response.numeric_value = (
                                decimal_or_none(
                                    numeric_value
                                )
                            )

                            # ------------------------------------
                            # NUMERATOR
                            # ------------------------------------

                            response.numerator = (
                                decimal_or_none(
                                    numerator
                                )
                            )

                            # ------------------------------------
                            # DENOMINATOR
                            # ------------------------------------

                            response.denominator = (
                                decimal_or_none(
                                    denominator
                                )
                            )

                            # ------------------------------------
                            # CONDUCTED
                            # ------------------------------------

                            if conducted == "yes":

                                response.conducted = True

                            elif conducted == "no":

                                response.conducted = False

                            else:

                                response.conducted = None

                            # ------------------------------------
                            # NOTES
                            # ------------------------------------

                            response.notes = notes

                            # ------------------------------------
                            # SAVE
                            # ------------------------------------

                            response.save()

                # =================================================
                # CALCULATE SCORE
                # =================================================

                calculate_assessment_total(
                    assessment
                )

                # =================================================
                # SAVE DRAFT
                # =================================================

                if action == "save":

                    messages.success(
                        request,
                        "Assessment draft saved successfully.",
                    )

                    return redirect(
                        "assessments:assessment_detail",
                        assessment_id=assessment.pk,
                    )

                # =================================================
                # VALIDATE
                # =================================================

                errors = validate_assessment(
                    assessment,
                    sections,
                )

                if errors:

                    messages.error(
                        request,
                        (
                            "The assessment cannot be submitted. "
                            f"{len(errors)} indicator(s) "
                            "are incomplete."
                        ),
                    )

                    for error in errors[:20]:

                        messages.warning(
                            request,
                            error,
                        )

                    return redirect(
                        "assessments:assessment_detail",
                        assessment_id=assessment.pk,
                    )

                # =================================================
                # SUBMIT
                # =================================================

                assessment.status = (
                    Assessment.Status.SUBMITTED
                )

                assessment.submitted_at = (
                    timezone.now()
                )

                if hasattr(
                    assessment,
                    "reviewed_at",
                ):

                    assessment.reviewed_at = None

                if hasattr(
                    assessment,
                    "approved_at",
                ):

                    assessment.approved_at = None

                assessment.save()

            messages.success(
                request,
                "Assessment submitted successfully.",
            )

        except Exception as exc:

            messages.error(
                request,
                (
                    "Unable to save the assessment: "
                    f"{exc}"
                ),
            )

        return redirect(
            "assessments:assessment_detail",
            assessment_id=assessment.pk,
        )

    # ============================================================
    # BUILD DISPLAY DATA
    # ============================================================

    response_forms = []

    for section in sections:

        for thematic_area in (
            section.thematic_areas
            .filter(active=True)
        ):

            indicators = (
                thematic_area.indicators
                .filter(active=True)
                .order_by(
                    "order",
                    "code",
                )
            )

            for indicator in indicators:

                response_forms.append(
                    {
                        "section": section,
                        "thematic_area": thematic_area,
                        "indicator": indicator,
                        "response": existing_responses.get(
                            indicator.id
                        ),
                    }
                )

    # ============================================================
    # VIEW ONLY
    # ============================================================

    is_submitted = assessment.status in [
        Assessment.Status.SUBMITTED,
        Assessment.Status.REVIEWED,
        Assessment.Status.APPROVED,
    ]

    return render(
        request,
        "assessments/assessment_detail.html",
        {
            "profile": profile,
            "assessment": assessment,
            "sections": sections,
            "response_forms": response_forms,
            "can_edit": can_edit,
            "can_review": can_review,
            "can_approve": can_approve,
            "is_submitted": is_submitted,
        },
    )


# ================================================================
# VALIDATE ASSESSMENT
# ================================================================

def validate_assessment(
    assessment,
    sections,
):
    """
    Validate every active indicator before submission.
    """

    errors = []

    responses = {
        response.indicator_id: response
        for response in (
            assessment.responses
            .select_related(
                "selected_option",
            )
        )
    }

    for section in sections:

        for thematic_area in (
            section.thematic_areas
            .filter(active=True)
        ):

            indicators = (
                thematic_area.indicators
                .filter(active=True)
                .prefetch_related("options")
                .order_by(
                    "order",
                    "code",
                )
            )

            for indicator in indicators:

                response = responses.get(
                    indicator.id
                )

                # =================================================
                # NO RESPONSE
                # =================================================

                if not response:

                    errors.append(
                        f"{indicator.code}: "
                        "No response has been entered."
                    )

                    continue

                # =================================================
                # CONDUCTED / NOT CONDUCTED
                # =================================================

                if (
                    indicator.response_type
                    == AssessmentIndicator.ResponseType.CONDUCTED_CHECK
                ):

                    if response.conducted is None:

                        errors.append(
                            f"{indicator.code}: "
                            "Please select Conducted or "
                            "Not Conducted."
                        )

                        continue

                    if response.conducted is False:
                        continue

                # =================================================
                # YES / NO
                # =================================================

                if (
                    indicator.response_type
                    == AssessmentIndicator.ResponseType.YES_NO
                ):

                    if not response.selected_option:

                        errors.append(
                            f"{indicator.code}: "
                            "Please select Yes or No."
                        )

                        continue

                    if (
                        response.selected_option.indicator_id
                        != indicator.id
                    ):

                        errors.append(
                            f"{indicator.code}: "
                            "Invalid answer selected."
                        )

                        continue

                # =================================================
                # MULTIPLE CHOICE
                # =================================================

                if (
                    indicator.response_type
                    == AssessmentIndicator.ResponseType.CHOICE
                ):

                    if not response.selected_option:

                        errors.append(
                            f"{indicator.code}: "
                            "Please select an answer."
                        )

                        continue

                    if (
                        response.selected_option.indicator_id
                        != indicator.id
                    ):

                        errors.append(
                            f"{indicator.code}: "
                            "Invalid answer selected."
                        )

                        continue

                # =================================================
                # NUMERIC
                # =================================================

                if (
                    indicator.response_type
                    == AssessmentIndicator.ResponseType.NUMERIC
                ):

                    if response.numeric_value is None:

                        errors.append(
                            f"{indicator.code}: "
                            "Numeric value is required."
                        )

                        continue

                    if response.numeric_value < 0:

                        errors.append(
                            f"{indicator.code}: "
                            "Numeric value cannot be negative."
                        )

                        continue

                # =================================================
                # RATIO
                # =================================================

                ratio_not_conducted = (
                    indicator.response_type
                    == AssessmentIndicator.ResponseType.CONDUCTED_CHECK
                    and response.conducted is False
                )

                if (
                    indicator.calculation_type
                    == AssessmentIndicator.CalculationType.RATIO
                    and not ratio_not_conducted
                ):

                    if response.numerator is None:

                        errors.append(
                            f"{indicator.code}: "
                            "Numerator is required."
                        )

                    elif response.numerator < 0:

                        errors.append(
                            f"{indicator.code}: "
                            "Numerator cannot be negative."
                        )

                    if response.denominator is None:

                        errors.append(
                            f"{indicator.code}: "
                            "Denominator is required."
                        )

                    elif response.denominator <= 0:

                        errors.append(
                            f"{indicator.code}: "
                            "Denominator must be greater than zero."
                        )

                    if (
                        response.numerator is not None
                        and response.denominator is not None
                        and response.denominator > 0
                        and response.numerator
                        > response.denominator
                    ):

                        errors.append(
                            f"{indicator.code}: "
                            "Numerator cannot be greater "
                            "than denominator."
                        )

    return errors


# ================================================================
# CALCULATE TOTAL SCORE
# ================================================================

def calculate_assessment_total(assessment):
    """
    Calculate the weighted assessment score.

    Expected weights:

        Health Infrastructure = 30%
        Data Quality          = 30%
        Data Use              = 40%

    Final score = 100%.
    """

    sections = (
        AssessmentSection.objects
        .filter(active=True)
        .prefetch_related(
            "thematic_areas__indicators"
        )
        .order_by(
            "order",
            "code",
        )
    )

    responses = {
        response.indicator_id: response
        for response in assessment.responses.all()
    }

    total = Decimal("0")

    for section in sections:

        maximum = Decimal("0")
        achieved = Decimal("0")

        for thematic_area in section.thematic_areas.all():

            for indicator in (
                thematic_area.indicators
                .filter(active=True)
            ):

                maximum_indicator = (
                    indicator.maximum_score
                    or Decimal("0")
                )

                maximum += maximum_indicator

                response = responses.get(
                    indicator.id
                )

                if not response:
                    continue

                if response.conducted is False:

                    score = Decimal("0")

                else:

                    score = (
                        response.calculated_score
                        or Decimal("0")
                    )

                if score > maximum_indicator:

                    score = maximum_indicator

                if score < 0:

                    score = Decimal("0")

                achieved += score

        if maximum <= 0:
            continue

        if achieved > maximum:

            achieved = maximum

        section_percentage = (
            achieved
            / maximum
        ) * Decimal("100")

        weighted_score = (
            section_percentage
            * (
                section.weight
                / Decimal("100")
            )
        )

        total += weighted_score

    total = total.quantize(
        Decimal("0.01")
    )

    assessment.total_score = total

    assessment.save(
        update_fields=[
            "total_score",
            "updated_at",
        ]
    )

    return total


# ================================================================
# ASSESSMENT DASHBOARD
# ================================================================

@login_required
def dashboard(request):
    """
    Assessment monitoring dashboard.

    Facility:
        Own facility.

    Sub-city:
        Own sub-city and its health centers.

    Region:
        Region -> sub-cities -> health centers.

    Ministry:
        Region -> sub-cities -> facilities.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    facilities = get_accessible_facilities(
        profile
    )

    regions = get_accessible_regions(
        profile
    )

    # ============================================================
    # REGION SELECTION
    # ============================================================

    selected_region = None

    region_id = request.GET.get(
        "region"
    )

    if profile.is_region_admin():

        selected_region = profile.organization

    elif region_id:

        selected_region = (
            regions
            .filter(
                id=region_id
            )
            .first()
        )

    # ============================================================
    # SUB-CITY SELECTION
    # ============================================================

    subcities = get_accessible_subcities(
        profile,
        selected_region,
    )

    selected_subcity = None

    subcity_id = request.GET.get(
        "subcity"
    )

    if profile.is_subcity_admin():

        selected_subcity = profile.organization

    elif subcity_id:

        selected_subcity = (
            subcities
            .filter(
                id=subcity_id
            )
            .first()
        )

    # ============================================================
    # FACILITY FILTERING
    # ============================================================

    if profile.is_facility_user():

        facilities = facilities.filter(
            id=profile.facility_id
        )

    elif selected_subcity:

        facilities = facilities.filter(
            organization=selected_subcity
        )

    elif selected_region:

        facilities = (
            facilities.filter(
                organization=selected_region
            )
            |
            facilities.filter(
                organization__parent=selected_region
            )
        ).distinct()

    # ============================================================
    # PERIODS
    # ============================================================

    periods = get_active_periods()

    years = get_years()

    selected_year = request.GET.get(
        "year"
    )

    try:

        selected_year = (
            int(selected_year)
            if selected_year
            else None
        )

    except (
        TypeError,
        ValueError,
    ):

        selected_year = None

    if (
        selected_year
        and selected_year not in years
    ):

        selected_year = None

    available_periods = periods

    if selected_year:

        available_periods = (
            available_periods
            .filter(
                year=selected_year
            )
        )

    # ============================================================
    # PERIOD SELECTION
    # ============================================================

    period_id = request.GET.get(
        "period"
    )

    selected_period = None

    if period_id:

        selected_period = (
            available_periods
            .filter(
                id=period_id
            )
            .first()
        )

    if not selected_period:

        selected_period = (
            available_periods
            .order_by(
                "-year",
                "-start_month",
                "-period_number",
            )
            .first()
        )

    # ============================================================
    # FACILITY USER
    # ============================================================

    if profile.is_facility_user():

        assessment_queryset = (
            Assessment.objects
            .filter(
                facility_id=profile.facility_id
            )
        )

    else:

        assessment_queryset = (
            Assessment.objects
            .filter(
                facility__in=facilities
            )
        )

    if selected_period:

        assessment_queryset = (
            assessment_queryset
            .filter(
                period=selected_period
            )
        )

    assessments = (
        assessment_queryset
        .select_related(
            "facility",
            "facility__organization",
            "period",
        )
        .prefetch_related(
            "responses"
        )
        .order_by(
            "facility__organization__name",
            "facility__name",
        )
    )

    assessment_map = {
        assessment.facility_id: assessment
        for assessment in assessments
    }

    # ============================================================
    # FACILITY ROWS
    # ============================================================

    facility_rows = []

    for facility in facilities:

        assessment = assessment_map.get(
            facility.id
        )

        if assessment:

            status = assessment.status

            status_label = (
                assessment.get_status_display()
            )

            score = assessment.total_score

        else:

            status = "NOT_STARTED"

            status_label = "Not Started"

            score = None

        region = get_region_for_facility(
            facility
        )

        subcity = get_subcity_for_facility(
            facility
        )

        facility_rows.append(
            {
                "facility": facility,
                "assessment": assessment,
                "status": status,
                "status_label": status_label,
                "score": score,
                "region": region,
                "subcity": subcity,
            }
        )

    # ============================================================
    # SUB-CITY GROUPS
    # ============================================================

    subcity_groups = []

    for subcity in subcities:

        group_rows = [
            row
            for row in facility_rows
            if (
                row["subcity"]
                and row["subcity"].id
                == subcity.id
            )
        ]

        subcity_groups.append(
            {
                "subcity": subcity,
                "facilities": group_rows,
            }
        )

    # ============================================================
    # COUNTS
    # ============================================================

    total_facilities = len(
        facility_rows
    )

    submitted_count = sum(
        1
        for row in facility_rows
        if row["status"]
        == Assessment.Status.SUBMITTED
    )

    reviewed_count = sum(
        1
        for row in facility_rows
        if row["status"]
        == Assessment.Status.REVIEWED
    )

    approved_count = sum(
        1
        for row in facility_rows
        if row["status"]
        == Assessment.Status.APPROVED
    )

    draft_count = sum(
        1
        for row in facility_rows
        if row["status"]
        == Assessment.Status.DRAFT
    )

    returned_count = sum(
        1
        for row in facility_rows
        if row["status"]
        == Assessment.Status.RETURNED
    )

    not_started_count = sum(
        1
        for row in facility_rows
        if row["status"]
        == "NOT_STARTED"
    )

    # ============================================================
    # CONTEXT
    # ============================================================

    context = {
        "profile": profile,

        "regions": regions,
        "subcities": subcities,
        "facilities": facilities,

        "selected_region": selected_region,
        "selected_subcity": selected_subcity,

        "years": years,
        "periods": available_periods,
        "selected_year": selected_year,
        "selected_period": selected_period,

        "facility_rows": facility_rows,
        "subcity_groups": subcity_groups,

        "total_facilities": total_facilities,
        "submitted_count": submitted_count,
        "reviewed_count": reviewed_count,
        "approved_count": approved_count,
        "draft_count": draft_count,
        "returned_count": returned_count,
        "not_started_count": not_started_count,

        "is_ministry": profile.is_ministry_admin(),
        "is_region": profile.is_region_admin(),
        "is_subcity": profile.is_subcity_admin(),
        "is_facility": profile.is_facility_user(),
    }

    return render(
        request,
        "assessments/assessment_dashboard.html",
        context,
    )


# ================================================================
# REVIEW
# ================================================================

@login_required
def assessment_review(
    request,
    assessment_id,
):
    """
    Review a submitted assessment.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    assessment = get_object_or_404(
        Assessment.objects.select_related(
            "facility",
            "facility__organization",
            "period",
            "created_by",
        ),
        pk=assessment_id,
    )

    if not user_can_review_assessment(
        profile,
        assessment,
    ):

        messages.error(
            request,
            "You do not have permission to review this assessment.",
        )

        return redirect(
            "assessments:dashboard"
        )

    if assessment.status != (
        Assessment.Status.SUBMITTED
    ):

        messages.warning(
            request,
            "Only submitted assessments can be reviewed.",
        )

        return redirect(
            "assessments:assessment_detail",
            assessment_id=assessment.pk,
        )

    if request.method == "POST":

        assessment.status = (
            Assessment.Status.REVIEWED
        )

        assessment.reviewed_at = (
            timezone.now()
        )

        assessment.save()

        messages.success(
            request,
            "Assessment reviewed successfully.",
        )

        return redirect(
            "assessments:assessment_detail",
            assessment_id=assessment.pk,
        )

    return render(
        request,
        "assessments/assessment_review.html",
        {
            "profile": profile,
            "assessment": assessment,
        },
    )


# ================================================================
# APPROVE
# ================================================================

@login_required
def assessment_approve(
    request,
    assessment_id,
):
    """
    Approve a reviewed assessment.

    Only Ministry and Region users can approve.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    assessment = get_object_or_404(
        Assessment.objects.select_related(
            "facility",
            "facility__organization",
            "period",
        ),
        pk=assessment_id,
    )

    if not user_can_approve_assessment(
        profile,
        assessment,
    ):

        messages.error(
            request,
            "You do not have permission to approve this assessment.",
        )

        return redirect(
            "assessments:dashboard"
        )

    if assessment.status != (
        Assessment.Status.REVIEWED
    ):

        messages.warning(
            request,
            "Only reviewed assessments can be approved.",
        )

        return redirect(
            "assessments:assessment_detail",
            assessment_id=assessment.pk,
        )

    if request.method != "POST":

        return redirect(
            "assessments:assessment_detail",
            assessment_id=assessment.pk,
        )

    assessment.status = (
        Assessment.Status.APPROVED
    )

    assessment.approved_at = (
        timezone.now()
    )

    assessment.save()

    messages.success(
        request,
        "Assessment approved successfully.",
    )

    return redirect(
        "assessments:assessment_detail",
        assessment_id=assessment.pk,
    )


# ================================================================
# SUCCESS
# ================================================================

@login_required
def assessment_success(
    request,
    assessment_id,
):
    """
    Assessment submission confirmation.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    assessment = get_object_or_404(
        Assessment.objects.select_related(
            "facility",
            "facility__organization",
            "period",
        ),
        pk=assessment_id,
    )

    if not user_can_access_assessment(
        profile,
        assessment,
    ):

        messages.error(
            request,
            "You do not have permission to view this assessment.",
        )

        return redirect(
            "assessments:dashboard"
        )

    return render(
        request,
        "assessments/success.html",
        {
            "profile": profile,
            "assessment": assessment,
        },
    )


# ################################################################
# ################################################################
#
#                    EXPORT FUNCTIONS
#
# ################################################################
# ################################################################


# ================================================================
# EXPORT DATA HELPER
# ================================================================

def _get_export_data(assessment):
    """
    Build the exact hierarchical data required by PDF and Excel.

    This uses the same active sections, thematic areas and
    indicators used by assessment_detail().
    """

    sections = (
        AssessmentSection.objects
        .filter(active=True)
        .prefetch_related(
            "thematic_areas__indicators__options"
        )
        .order_by(
            "order",
            "code",
        )
    )

    responses = {
        response.indicator_id: response
        for response in (
            assessment.responses
            .select_related(
                "selected_option",
                "indicator",
            )
        )
    }

    export_sections = []

    for section in sections:

        section_maximum = Decimal("0")
        section_achieved = Decimal("0")

        thematic_rows = []

        for thematic_area in (
            section.thematic_areas
            .filter(active=True)
        ):

            thematic_maximum = Decimal("0")
            thematic_achieved = Decimal("0")

            indicator_rows = []

            indicators = (
                thematic_area.indicators
                .filter(active=True)
                .prefetch_related("options")
                .order_by(
                    "order",
                    "code",
                )
            )

            for indicator in indicators:

                response = responses.get(
                    indicator.id
                )

                maximum_score = (
                    indicator.maximum_score
                    or Decimal("0")
                )

                score = Decimal("0")

                if response:

                    if response.conducted is False:

                        score = Decimal("0")

                    else:

                        score = (
                            response.calculated_score
                            or Decimal("0")
                        )

                if score < 0:
                    score = Decimal("0")

                if score > maximum_score:
                    score = maximum_score

                section_maximum += maximum_score
                section_achieved += score

                thematic_maximum += maximum_score
                thematic_achieved += score

                # ------------------------------------------------
                # ANSWER
                # ------------------------------------------------

                answer_parts = []

                selected_option = None

                if response:

                    selected_option = (
                        response.selected_option
                    )

                if selected_option:

                    answer_parts.append(
                        str(
                            getattr(
                                selected_option,
                                "label",
                                selected_option,
                            )
                        )
                    )

                if (
                    response
                    and response.numeric_value is not None
                ):

                    answer_parts.append(
                        f"Numeric: {response.numeric_value}"
                    )

                if response:

                    if (
                        response.numerator is not None
                        or response.denominator is not None
                    ):

                        numerator_text = (
                            response.numerator
                            if response.numerator is not None
                            else "—"
                        )

                        denominator_text = (
                            response.denominator
                            if response.denominator is not None
                            else "—"
                        )

                        answer_parts.append(
                            f"Numerator: {numerator_text} / "
                            f"Denominator: {denominator_text}"
                        )

                if (
                    response
                    and response.conducted is not None
                ):

                    if response.conducted:

                        answer_parts.append(
                            "Conducted"
                        )

                    else:

                        answer_parts.append(
                            "Not Conducted"
                        )

                answer = (
                    " | ".join(answer_parts)
                    if answer_parts
                    else "No response"
                )

                # ------------------------------------------------
                # RATIO
                # ------------------------------------------------

                ratio_percentage = None

                if (
                    response
                    and response.numerator is not None
                    and response.denominator is not None
                    and response.denominator > 0
                ):

                    ratio_percentage = (
                        response.numerator
                        / response.denominator
                    ) * Decimal("100")

                    if ratio_percentage > Decimal("100"):
                        ratio_percentage = Decimal("100")

                    if ratio_percentage < Decimal("0"):
                        ratio_percentage = Decimal("0")

                # ------------------------------------------------
                # NOTES
                # ------------------------------------------------

                notes = ""

                if response:

                    notes = getattr(
                        response,
                        "notes",
                        "",
                    ) or ""

                indicator_rows.append(
                    {
                        "indicator": indicator,
                        "response": response,
                        "answer": answer,
                        "maximum_score": maximum_score,
                        "score": score,
                        "ratio_percentage": ratio_percentage,
                        "notes": notes,
                    }
                )

            thematic_percentage = Decimal("0")

            if thematic_maximum > 0:

                thematic_percentage = (
                    thematic_achieved
                    / thematic_maximum
                ) * Decimal("100")

            thematic_rows.append(
                {
                    "thematic_area": thematic_area,
                    "indicators": indicator_rows,
                    "maximum": thematic_maximum,
                    "achieved": thematic_achieved,
                    "percentage": thematic_percentage,
                }
            )

        section_percentage = Decimal("0")

        if section_maximum > 0:

            section_percentage = (
                section_achieved
                / section_maximum
            ) * Decimal("100")

        weighted_score = (
            section_percentage
            * (
                section.weight
                / Decimal("100")
            )
        )

        export_sections.append(
            {
                "section": section,
                "thematics": thematic_rows,
                "maximum": section_maximum,
                "achieved": section_achieved,
                "percentage": section_percentage,
                "weighted_score": weighted_score,
            }
        )

    return export_sections


# ================================================================
# EXPORT VALUE FORMATTER
# ================================================================

def _export_decimal(value, places=2):
    """
    Convert Decimal / numeric values to a clean string.
    """

    if value is None:
        return ""

    try:

        decimal_value = Decimal(str(value))

        return f"{decimal_value:.{places}f}"

    except (
        InvalidOperation,
        TypeError,
        ValueError,
    ):

        return str(value)


# ================================================================
# PDF EXPORT
# ================================================================

@login_required
def export_assessment_pdf(
    request,
    assessment_id,
):
    """
    Export one assessment as a PDF.

    Access is exactly the same as assessment_detail().
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    assessment = get_object_or_404(
        Assessment.objects.select_related(
            "facility",
            "facility__organization",
            "facility__organization__parent",
            "period",
            "created_by",
        ),
        pk=assessment_id,
    )

    # ============================================================
    # ACCESS CONTROL
    # ============================================================

    if not user_can_access_assessment(
        profile,
        assessment,
    ):

        messages.error(
            request,
            "You do not have permission to export this assessment.",
        )

        return redirect(
            "assessments:dashboard"
        )

    export_sections = _get_export_data(
        assessment
    )

    # ============================================================
    # ORGANIZATION INFORMATION
    # ============================================================

    facility = assessment.facility

    organization = getattr(
        facility,
        "organization",
        None,
    )

    region = get_region_for_facility(
        facility
    )

    subcity = get_subcity_for_facility(
        facility
    )

    # ============================================================
    # PDF BUFFER
    # ============================================================

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=(
            f"EHIRS Assessment - "
            f"{facility.name}"
        ),
        author="EHIRS",
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "EHIRSTitle",
        parent=styles["Title"],
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#115e59"),
        spaceAfter=5,
    )

    subtitle_style = ParagraphStyle(
        "EHIRSSubtitle",
        parent=styles["Normal"],
        fontSize=9,
        leading=12,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=12,
    )

    heading_style = ParagraphStyle(
        "EHIRSHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#115e59"),
        spaceBefore=8,
        spaceAfter=6,
    )

    small_style = ParagraphStyle(
        "EHIRSSmall",
        parent=styles["Normal"],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#334155"),
    )

    small_center_style = ParagraphStyle(
        "EHIRSSmallCenter",
        parent=small_style,
        alignment=TA_CENTER,
    )

    normal_style = ParagraphStyle(
        "EHIRSNormal",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor("#1e293b"),
    )

    story = []

    # ============================================================
    # TITLE
    # ============================================================

    story.append(
        Paragraph(
            "EHIRS HEALTH FACILITY ASSESSMENT",
            title_style,
        )
    )

    story.append(
        Paragraph(
            "Early Health Information Reporting System",
            subtitle_style,
        )
    )

    # ============================================================
    # SUMMARY
    # ============================================================

    status_label = (
        assessment.get_status_display()
    )

    facility_name = getattr(
        facility,
        "name",
        "—",
    )

    organization_name = (
        getattr(
            organization,
            "name",
            "—",
        )
        if organization
        else "—"
    )

    region_name = (
        getattr(
            region,
            "name",
            "—",
        )
        if region
        else "—"
    )

    subcity_name = (
        getattr(
            subcity,
            "name",
            "—",
        )
        if subcity
        else "—"
    )

    period_name = getattr(
        assessment.period,
        "name",
        "—",
    )

    period_year = getattr(
        assessment.period,
        "year",
        "—",
    )

    summary_data = [
        [
            Paragraph("<b>Facility</b>", small_style),
            Paragraph(
                str(facility_name),
                small_style,
            ),
            Paragraph("<b>Organization</b>", small_style),
            Paragraph(
                str(organization_name),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Region</b>", small_style),
            Paragraph(
                str(region_name),
                small_style,
            ),
            Paragraph("<b>Sub-city</b>", small_style),
            Paragraph(
                str(subcity_name),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Year</b>", small_style),
            Paragraph(
                str(period_year),
                small_style,
            ),
            Paragraph("<b>Period</b>", small_style),
            Paragraph(
                str(period_name),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Assessment ID</b>", small_style),
            Paragraph(
                str(assessment.id),
                small_style,
            ),
            Paragraph("<b>Status</b>", small_style),
            Paragraph(
                str(status_label),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Overall Score</b>", small_style),
            Paragraph(
                f"{_export_decimal(assessment.total_score)} / 100",
                small_style,
            ),
            Paragraph("<b>Created By</b>", small_style),
            Paragraph(
                str(
                    getattr(
                        assessment.created_by,
                        "get_username",
                        lambda: "—",
                    )()
                    if assessment.created_by
                    else "—"
                ),
                small_style,
            ),
        ],
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            32 * mm,
            75 * mm,
            32 * mm,
            75 * mm,
        ],
        repeatRows=0,
    )

    summary_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#f1f5f9"),
                ),
                (
                    "BACKGROUND",
                    (2, 0),
                    (2, -1),
                    colors.HexColor("#f1f5f9"),
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.7,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "INNERGRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#dbe4e8"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    story.append(summary_table)
    story.append(Spacer(1, 7 * mm))

    # ============================================================
    # SECTION SUMMARY
    # ============================================================

    story.append(
        Paragraph(
            "Section Summary",
            heading_style,
        )
    )

    section_summary_data = [
        [
            Paragraph("<b>Section</b>", small_center_style),
            Paragraph("<b>Name</b>", small_center_style),
            Paragraph("<b>Weight</b>", small_center_style),
            Paragraph("<b>Maximum</b>", small_center_style),
            Paragraph("<b>Achieved</b>", small_center_style),
            Paragraph("<b>Percentage</b>", small_center_style),
            Paragraph("<b>Weighted Score</b>", small_center_style),
        ]
    ]

    for item in export_sections:

        section = item["section"]

        section_summary_data.append(
            [
                Paragraph(
                    str(section.code),
                    small_center_style,
                ),
                Paragraph(
                    str(section.name),
                    small_style,
                ),
                Paragraph(
                    f"{_export_decimal(section.weight)}%",
                    small_center_style,
                ),
                Paragraph(
                    _export_decimal(item["maximum"]),
                    small_center_style,
                ),
                Paragraph(
                    _export_decimal(item["achieved"]),
                    small_center_style,
                ),
                Paragraph(
                    f"{_export_decimal(item['percentage'])}%",
                    small_center_style,
                ),
                Paragraph(
                    _export_decimal(
                        item["weighted_score"]
                    ),
                    small_center_style,
                ),
            ]
        )

    section_table = Table(
        section_summary_data,
        colWidths=[
            20 * mm,
            72 * mm,
            24 * mm,
            28 * mm,
            28 * mm,
            32 * mm,
            32 * mm,
        ],
        repeatRows=1,
    )

    section_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#115e59"),
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.white,
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.7,
                    colors.HexColor("#94a3b8"),
                ),
                (
                    "INNERGRID",
                    (0, 0),
                    (-1, -1),
                    0.35,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    4,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    5,
                ),
            ]
        )
    )

    story.append(section_table)

    # ============================================================
    # CHECKLIST DETAILS
    # ============================================================

    for section_index, section_item in enumerate(
        export_sections,
        start=1,
    ):

        section = section_item["section"]

        story.append(
            PageBreak()
        )

        story.append(
            Paragraph(
                (
                    f"{section_index}. "
                    f"{section.code} - {section.name}"
                ),
                heading_style,
            )
        )

        section_info = (
            f"Weight: "
            f"{_export_decimal(section.weight)}% "
            f"&nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Score: "
            f"{_export_decimal(section_item['achieved'])} / "
            f"{_export_decimal(section_item['maximum'])} "
            f"&nbsp;&nbsp;|&nbsp;&nbsp; "
            f"Performance: "
            f"{_export_decimal(section_item['percentage'])}%"
        )

        story.append(
            Paragraph(
                section_info,
                small_style,
            )
        )

        story.append(
            Spacer(1, 4 * mm)
        )

        for thematic_item in section_item["thematics"]:

            thematic = thematic_item[
                "thematic_area"
            ]

            story.append(
                Paragraph(
                    (
                        f"{getattr(thematic, 'code', '')} "
                        f"- {getattr(thematic, 'name', '')}"
                    ),
                    ParagraphStyle(
                        "ThematicPDF",
                        parent=heading_style,
                        fontSize=10,
                        textColor=colors.HexColor(
                            "#0369a1"
                        ),
                        spaceBefore=5,
                        spaceAfter=4,
                    ),
                )
            )

            indicator_data = [
                [
                    Paragraph("<b>Code</b>", small_center_style),
                    Paragraph("<b>Indicator / Question</b>", small_style),
                    Paragraph("<b>Response</b>", small_style),
                    Paragraph("<b>Numeric</b>", small_center_style),
                    Paragraph("<b>Numerator</b>", small_center_style),
                    Paragraph("<b>Denominator</b>", small_center_style),
                    Paragraph("<b>Ratio %</b>", small_center_style),
                    Paragraph("<b>Score</b>", small_center_style),
                    Paragraph("<b>Maximum</b>", small_center_style),
                    Paragraph("<b>Notes</b>", small_style),
                ]
            ]

            for row in thematic_item["indicators"]:

                indicator = row["indicator"]
                response = row["response"]

                numeric_value = ""

                numerator_value = ""
                denominator_value = ""

                conducted_value = ""

                if response:

                    if response.numeric_value is not None:

                        numeric_value = _export_decimal(
                            response.numeric_value
                        )

                    if response.numerator is not None:

                        numerator_value = _export_decimal(
                            response.numerator
                        )

                    if response.denominator is not None:

                        denominator_value = _export_decimal(
                            response.denominator
                        )

                    if response.conducted is not None:

                        conducted_value = (
                            "Conducted"
                            if response.conducted
                            else "Not Conducted"
                        )

                answer_text = row["answer"]

                if conducted_value and (
                    conducted_value not in answer_text
                ):

                    if answer_text == "No response":

                        answer_text = conducted_value

                    else:

                        answer_text = (
                            f"{answer_text} | "
                            f"{conducted_value}"
                        )

                ratio_text = ""

                if (
                    row["ratio_percentage"]
                    is not None
                ):

                    ratio_text = (
                        _export_decimal(
                            row["ratio_percentage"]
                        )
                        + "%"
                    )

                indicator_data.append(
                    [
                        Paragraph(
                            str(
                                getattr(
                                    indicator,
                                    "code",
                                    "",
                                )
                            ),
                            small_center_style,
                        ),
                        Paragraph(
                            str(
                                getattr(
                                    indicator,
                                    "question",
                                    "",
                                )
                            ),
                            small_style,
                        ),
                        Paragraph(
                            str(answer_text),
                            small_style,
                        ),
                        Paragraph(
                            str(numeric_value),
                            small_center_style,
                        ),
                        Paragraph(
                            str(numerator_value),
                            small_center_style,
                        ),
                        Paragraph(
                            str(denominator_value),
                            small_center_style,
                        ),
                        Paragraph(
                            str(ratio_text),
                            small_center_style,
                        ),
                        Paragraph(
                            _export_decimal(
                                row["score"]
                            ),
                            small_center_style,
                        ),
                        Paragraph(
                            _export_decimal(
                                row["maximum_score"]
                            ),
                            small_center_style,
                        ),
                        Paragraph(
                            str(row["notes"]),
                            small_style,
                        ),
                    ]
                )

            indicator_table = Table(
                indicator_data,
                colWidths=[
                    17 * mm,
                    57 * mm,
                    42 * mm,
                    20 * mm,
                    20 * mm,
                    20 * mm,
                    19 * mm,
                    18 * mm,
                    20 * mm,
                    37 * mm,
                ],
                repeatRows=1,
            )

            indicator_table.setStyle(
                TableStyle(
                    [
                        (
                            "BACKGROUND",
                            (0, 0),
                            (-1, 0),
                            colors.HexColor("#0f766e"),
                        ),
                        (
                            "TEXTCOLOR",
                            (0, 0),
                            (-1, 0),
                            colors.white,
                        ),
                        (
                            "BOX",
                            (0, 0),
                            (-1, -1),
                            0.5,
                            colors.HexColor("#94a3b8"),
                        ),
                        (
                            "INNERGRID",
                            (0, 0),
                            (-1, -1),
                            0.3,
                            colors.HexColor("#cbd5e1"),
                        ),
                        (
                            "VALIGN",
                            (0, 0),
                            (-1, -1),
                            "TOP",
                        ),
                        (
                            "LEFTPADDING",
                            (0, 0),
                            (-1, -1),
                            3,
                        ),
                        (
                            "RIGHTPADDING",
                            (0, 0),
                            (-1, -1),
                            3,
                        ),
                        (
                            "TOPPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                        (
                            "BOTTOMPADDING",
                            (0, 0),
                            (-1, -1),
                            4,
                        ),
                        (
                            "ROWBACKGROUNDS",
                            (0, 1),
                            (-1, -1),
                            [
                                colors.white,
                                colors.HexColor("#f8fafc"),
                            ],
                        ),
                    ]
                )
            )

            story.append(
                indicator_table
            )

            story.append(
                Spacer(1, 4 * mm)
            )

    # ============================================================
    # WORKFLOW INFORMATION
    # ============================================================

    story.append(
        PageBreak()
    )

    story.append(
        Paragraph(
            "Assessment Workflow Information",
            heading_style,
        )
    )

    workflow_rows = [
        [
            Paragraph("<b>Status</b>", small_style),
            Paragraph(
                str(status_label),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Submitted At</b>", small_style),
            Paragraph(
                str(
                    getattr(
                        assessment,
                        "submitted_at",
                        None,
                    )
                    or "—"
                ),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Reviewed At</b>", small_style),
            Paragraph(
                str(
                    getattr(
                        assessment,
                        "reviewed_at",
                        None,
                    )
                    or "—"
                ),
                small_style,
            ),
        ],
        [
            Paragraph("<b>Approved At</b>", small_style),
            Paragraph(
                str(
                    getattr(
                        assessment,
                        "approved_at",
                        None,
                    )
                    or "—"
                ),
                small_style,
            ),
        ],
    ]

    workflow_table = Table(
        workflow_rows,
        colWidths=[
            45 * mm,
            150 * mm,
        ],
    )

    workflow_table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (0, -1),
                    colors.HexColor("#f1f5f9"),
                ),
                (
                    "BOX",
                    (0, 0),
                    (-1, -1),
                    0.6,
                    colors.HexColor("#cbd5e1"),
                ),
                (
                    "INNERGRID",
                    (0, 0),
                    (-1, -1),
                    0.3,
                    colors.HexColor("#dbe4e8"),
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE",
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    6,
                ),
            ]
        )
    )

    story.append(
        workflow_table
    )

    story.append(
        Spacer(1, 10 * mm)
    )

    story.append(
        Paragraph(
            (
                "Generated by EHIRS. "
                "This document represents the assessment data "
                "available in the system at the time of export."
            ),
            subtitle_style,
        )
    )

    # ============================================================
    # BUILD PDF
    # ============================================================

    document.build(
        story
    )

    pdf = buffer.getvalue()

    buffer.close()

    filename = (
        f"EHIRS_Assessment_"
        f"{facility.name}_"
        f"{assessment.id}.pdf"
    )

    # Clean filename
    filename = (
        filename
        .replace("/", "_")
        .replace("\\", "_")
        .replace(" ", "_")
    )

    response = HttpResponse(
        pdf,
        content_type="application/pdf",
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response

# ================================================================
# EXCEL EXPORT
# ================================================================

@login_required
def export_assessment_excel(
    request,
    assessment_id,
):
    """
    Export one assessment as an Excel workbook.

    Workbook sheets:

        1. Assessment Summary
        2. Checklist Details
        3. Section Summary
        4. Thematic Summary

    Important:
        Django stores DateTimeField values as timezone-aware
        datetimes when USE_TZ=True.

        Excel/openpyxl does not support timezone-aware
        datetime objects.

        Therefore all assessment timestamps are converted
        to readable text before being written to Excel.
    """

    profile = get_active_profile(request)

    if not profile:

        messages.error(
            request,
            "Your EHIRS user profile is inactive or unavailable.",
        )

        return redirect(
            "accounts:login"
        )

    assessment = get_object_or_404(
        Assessment.objects.select_related(
            "facility",
            "facility__organization",
            "period",
            "created_by",
        ),
        pk=assessment_id,
    )

    # ============================================================
    # ACCESS CONTROL
    # ============================================================

    if not user_can_access_assessment(
        profile,
        assessment,
    ):

        messages.error(
            request,
            "You do not have permission to export this assessment.",
        )

        return redirect(
            "assessments:dashboard"
        )

    # ============================================================
    # SAFE DATETIME DISPLAY
    # ============================================================

    def excel_datetime_text(value):
        """
        Convert a Django timezone-aware datetime to plain text.

        Excel will receive a string, not a timezone-aware
        datetime object.
        """

        if value is None:
            return ""

        try:

            if timezone.is_aware(value):

                value = timezone.localtime(
                    value
                )

            return value.strftime(
                "%Y-%m-%d %H:%M:%S"
            )

        except (
            AttributeError,
            TypeError,
            ValueError,
        ):

            return str(value)

    # ============================================================
    # EXPORT DATA
    # ============================================================

    export_sections = _get_export_data(
        assessment
    )

    facility = assessment.facility

    organization = getattr(
        facility,
        "organization",
        None,
    )

    region = get_region_for_facility(
        facility
    )

    subcity = get_subcity_for_facility(
        facility
    )

    # ============================================================
    # WORKBOOK
    # ============================================================

    workbook = Workbook()

    summary_sheet = workbook.active

    summary_sheet.title = (
        "Assessment Summary"
    )

    detail_sheet = workbook.create_sheet(
        "Checklist Details"
    )

    section_sheet = workbook.create_sheet(
        "Section Summary"
    )

    thematic_sheet = workbook.create_sheet(
        "Thematic Summary"
    )

    # ============================================================
    # STYLES
    # ============================================================

    primary_fill = PatternFill(
        fill_type="solid",
        fgColor="0F766E",
    )

    dark_fill = PatternFill(
        fill_type="solid",
        fgColor="115E59",
    )

    light_fill = PatternFill(
        fill_type="solid",
        fgColor="F0FDFA",
    )

    gray_fill = PatternFill(
        fill_type="solid",
        fgColor="F1F5F9",
    )

    blue_fill = PatternFill(
        fill_type="solid",
        fgColor="EFF6FF",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    title_font = Font(
        size=16,
        bold=True,
        color="115E59",
    )

    section_font = Font(
        size=12,
        bold=True,
        color="115E59",
    )

    bold_font = Font(
        bold=True,
    )

    thin_side = Side(
        style="thin",
        color="CBD5E1",
    )

    border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    center_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    left_alignment = Alignment(
        horizontal="left",
        vertical="top",
        wrap_text=True,
    )

    # ============================================================
    # SHEET 1 - ASSESSMENT SUMMARY
    # ============================================================

    summary_sheet.merge_cells(
        "A1:D1"
    )

    summary_sheet["A1"] = (
        "EHIRS HEALTH FACILITY ASSESSMENT"
    )

    summary_sheet["A1"].font = (
        title_font
    )

    summary_sheet["A1"].alignment = (
        center_alignment
    )

    summary_sheet.merge_cells(
        "A2:D2"
    )

    summary_sheet["A2"] = (
        "Early Health Information Reporting System"
    )

    summary_sheet["A2"].alignment = (
        center_alignment
    )

    # ============================================================
    # IMPORTANT:
    # All timestamp fields are converted to TEXT here.
    # They are NOT sent directly to Excel as Django datetimes.
    # ============================================================

    summary_rows = [

        (
            "Facility",
            getattr(
                facility,
                "name",
                "—",
            ),
        ),

        (
            "Organization",
            getattr(
                organization,
                "name",
                "—",
            )
            if organization
            else "—",
        ),

        (
            "Region",
            getattr(
                region,
                "name",
                "—",
            )
            if region
            else "—",
        ),

        (
            "Sub-city",
            getattr(
                subcity,
                "name",
                "—",
            )
            if subcity
            else "—",
        ),

        (
            "Assessment ID",
            assessment.id,
        ),

        (
            "Year",
            getattr(
                assessment.period,
                "year",
                "—",
            ),
        ),

        (
            "Period",
            getattr(
                assessment.period,
                "name",
                "—",
            ),
        ),

        (
            "Status",
            assessment.get_status_display(),
        ),

        (
            "Overall Score",
            float(
                assessment.total_score
                or Decimal("0")
            ),
        ),

        (
            "Created By",
            (
                assessment.created_by.username
                if assessment.created_by
                else "—"
            ),
        ),

        # --------------------------------------------------------
        # TIMEZONE-SAFE VALUES
        # --------------------------------------------------------

        (
            "Submitted At",
            excel_datetime_text(
                getattr(
                    assessment,
                    "submitted_at",
                    None,
                )
            ),
        ),

        (
            "Reviewed At",
            excel_datetime_text(
                getattr(
                    assessment,
                    "reviewed_at",
                    None,
                )
            ),
        ),

        (
            "Approved At",
            excel_datetime_text(
                getattr(
                    assessment,
                    "approved_at",
                    None,
                )
            ),
        ),
    ]

    row_number = 4

    for label, value in summary_rows:

        label_cell = summary_sheet.cell(
            row=row_number,
            column=1,
            value=label,
        )

        value_cell = summary_sheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        label_cell.font = bold_font
        label_cell.fill = gray_fill
        label_cell.border = border

        value_cell.border = border
        value_cell.alignment = (
            left_alignment
        )

        row_number += 1

    summary_sheet.column_dimensions[
        "A"
    ].width = 25

    summary_sheet.column_dimensions[
        "B"
    ].width = 45

    summary_sheet.column_dimensions[
        "C"
    ].width = 18

    summary_sheet.column_dimensions[
        "D"
    ].width = 18

    summary_sheet.freeze_panes = "A4"

    # ============================================================
    # SHEET 2 - CHECKLIST DETAILS
    # ============================================================

    detail_headers = [
        "Section Code",
        "Section",
        "Section Weight %",
        "Thematic Code",
        "Thematic Area",
        "Indicator Code",
        "Indicator / Question",
        "Response",
        "Numeric Value",
        "Numerator",
        "Denominator",
        "Ratio %",
        "Calculated Score",
        "Maximum Score",
        "Conducted",
        "Notes",
    ]

    for column, header in enumerate(
        detail_headers,
        start=1,
    ):

        cell = detail_sheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.fill = primary_fill
        cell.font = white_font
        cell.alignment = (
            center_alignment
        )
        cell.border = border

    detail_row = 2

    for section_item in export_sections:

        section = section_item[
            "section"
        ]

        for thematic_item in section_item[
            "thematics"
        ]:

            thematic = thematic_item[
                "thematic_area"
            ]

            for item in thematic_item[
                "indicators"
            ]:

                indicator = item[
                    "indicator"
                ]

                response = item[
                    "response"
                ]

                numeric_value = None
                numerator = None
                denominator = None
                conducted = ""

                if response:

                    numeric_value = (
                        response.numeric_value
                    )

                    numerator = (
                        response.numerator
                    )

                    denominator = (
                        response.denominator
                    )

                    if response.conducted is True:

                        conducted = (
                            "Conducted"
                        )

                    elif response.conducted is False:

                        conducted = (
                            "Not Conducted"
                        )

                ratio = (
                    float(
                        item[
                            "ratio_percentage"
                        ]
                    )
                    if item[
                        "ratio_percentage"
                    ] is not None
                    else None
                )

                values = [

                    getattr(
                        section,
                        "code",
                        "",
                    ),

                    getattr(
                        section,
                        "name",
                        "",
                    ),

                    float(
                        section.weight
                    ),

                    getattr(
                        thematic,
                        "code",
                        "",
                    ),

                    getattr(
                        thematic,
                        "name",
                        "",
                    ),

                    getattr(
                        indicator,
                        "code",
                        "",
                    ),

                    getattr(
                        indicator,
                        "question",
                        "",
                    ),

                    item[
                        "answer"
                    ],

                    (
                        float(
                            numeric_value
                        )
                        if numeric_value
                        is not None
                        else None
                    ),

                    (
                        float(
                            numerator
                        )
                        if numerator
                        is not None
                        else None
                    ),

                    (
                        float(
                            denominator
                        )
                        if denominator
                        is not None
                        else None
                    ),

                    ratio,

                    float(
                        item[
                            "score"
                        ]
                    ),

                    float(
                        item[
                            "maximum_score"
                        ]
                    ),

                    conducted,

                    item[
                        "notes"
                    ],
                ]

                for column, value in enumerate(
                    values,
                    start=1,
                ):

                    cell = detail_sheet.cell(
                        row=detail_row,
                        column=column,
                        value=value,
                    )

                    cell.border = border

                    if column in [
                        1,
                        3,
                        4,
                        6,
                        9,
                        10,
                        11,
                        12,
                        13,
                        14,
                    ]:

                        cell.alignment = (
                            center_alignment
                        )

                    else:

                        cell.alignment = (
                            left_alignment
                        )

                detail_row += 1

    # ============================================================
    # DETAIL COLUMN WIDTHS
    # ============================================================

    detail_widths = [
        15,
        28,
        16,
        16,
        28,
        18,
        55,
        42,
        16,
        14,
        14,
        12,
        17,
        15,
        18,
        35,
    ]

    for index, width in enumerate(
        detail_widths,
        start=1,
    ):

        detail_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    detail_sheet.freeze_panes = "A2"

    if detail_sheet.max_row >= 2:

        detail_sheet.auto_filter.ref = (
            detail_sheet.dimensions
        )

    # ============================================================
    # SHEET 3 - SECTION SUMMARY
    # ============================================================

    section_headers = [
        "Section Code",
        "Section Name",
        "Weight %",
        "Maximum Score",
        "Achieved Score",
        "Performance %",
        "Weighted Score",
    ]

    for column, header in enumerate(
        section_headers,
        start=1,
    ):

        cell = section_sheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.fill = dark_fill
        cell.font = white_font
        cell.alignment = (
            center_alignment
        )
        cell.border = border

    section_row = 2

    for item in export_sections:

        section = item[
            "section"
        ]

        values = [

            section.code,

            section.name,

            float(
                section.weight
            ),

            float(
                item[
                    "maximum"
                ]
            ),

            float(
                item[
                    "achieved"
                ]
            ),

            float(
                item[
                    "percentage"
                ]
            ),

            float(
                item[
                    "weighted_score"
                ]
            ),
        ]

        for column, value in enumerate(
            values,
            start=1,
        ):

            cell = section_sheet.cell(
                row=section_row,
                column=column,
                value=value,
            )

            cell.border = border

            if column in [
                3,
                4,
                5,
                6,
                7,
            ]:

                cell.alignment = (
                    center_alignment
                )

            else:

                cell.alignment = (
                    left_alignment
                )

        section_row += 1

    section_widths = [
        18,
        38,
        15,
        18,
        18,
        18,
        18,
    ]

    for index, width in enumerate(
        section_widths,
        start=1,
    ):

        section_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    section_sheet.freeze_panes = "A2"

    if section_sheet.max_row >= 2:

        section_sheet.auto_filter.ref = (
            section_sheet.dimensions
        )

    # ============================================================
    # SHEET 4 - THEMATIC SUMMARY
    # ============================================================

    thematic_headers = [
        "Section",
        "Thematic Code",
        "Thematic Area",
        "Maximum Score",
        "Achieved Score",
        "Performance %",
    ]

    for column, header in enumerate(
        thematic_headers,
        start=1,
    ):

        cell = thematic_sheet.cell(
            row=1,
            column=column,
            value=header,
        )

        cell.fill = primary_fill
        cell.font = white_font
        cell.alignment = (
            center_alignment
        )
        cell.border = border

    thematic_row = 2

    for section_item in export_sections:

        section = section_item[
            "section"
        ]

        for thematic_item in section_item[
            "thematics"
        ]:

            thematic = thematic_item[
                "thematic_area"
            ]

            values = [

                section.name,

                thematic.code,

                thematic.name,

                float(
                    thematic_item[
                        "maximum"
                    ]
                ),

                float(
                    thematic_item[
                        "achieved"
                    ]
                ),

                float(
                    thematic_item[
                        "percentage"
                    ]
                ),
            ]

            for column, value in enumerate(
                values,
                start=1,
            ):

                cell = thematic_sheet.cell(
                    row=thematic_row,
                    column=column,
                    value=value,
                )

                cell.border = border

                if column >= 4:

                    cell.alignment = (
                        center_alignment
                    )

                else:

                    cell.alignment = (
                        left_alignment
                    )

            thematic_row += 1

    thematic_widths = [
        35,
        18,
        40,
        20,
        20,
        20,
    ]

    for index, width in enumerate(
        thematic_widths,
        start=1,
    ):

        thematic_sheet.column_dimensions[
            get_column_letter(index)
        ].width = width

    thematic_sheet.freeze_panes = "A2"

    if thematic_sheet.max_row >= 2:

        thematic_sheet.auto_filter.ref = (
            thematic_sheet.dimensions
        )

    # ============================================================
    # GENERAL FORMATTING
    # ============================================================

    for sheet in workbook.worksheets:

        sheet.sheet_view.showGridLines = (
            False
        )

        for row in sheet.iter_rows():

            for cell in row:

                if cell.value is not None:

                    cell.alignment = Alignment(
                        horizontal=(
                            cell.alignment.horizontal
                            or "left"
                        ),
                        vertical="top",
                        wrap_text=True,
                    )

    # ============================================================
    # FINAL TIMEZONE SAFETY CHECK
    # ============================================================
    #
    # This is an additional safety layer.
    #
    # If ANY timezone-aware datetime or time object somehow
    # reaches a cell, remove its timezone before openpyxl saves
    # the workbook.
    #
    # ============================================================

    for sheet in workbook.worksheets:

        for row in sheet.iter_rows():

            for cell in row:

                value = cell.value

                # ----------------------------------------------
                # DATETIME
                # ----------------------------------------------

                if isinstance(
                    value,
                    datetime,
                ):

                    if value.tzinfo is not None:

                        cell.value = (
                            value.replace(
                                tzinfo=None
                            )
                        )

                # ----------------------------------------------
                # TIME
                # ----------------------------------------------

                elif isinstance(
                    value,
                    time,
                ):

                    if value.tzinfo is not None:

                        cell.value = (
                            value.replace(
                                tzinfo=None
                            )
                        )

    # ============================================================
    # SAVE WORKBOOK TO MEMORY
    # ============================================================

    buffer = BytesIO()

    workbook.save(
        buffer
    )

    buffer.seek(0)

    # ============================================================
    # FILE NAME
    # ============================================================

    filename = (
        f"EHIRS_Assessment_"
        f"{facility.name}_"
        f"{assessment.id}.xlsx"
    )

    filename = (
        filename
        .replace(
            "/",
            "_",
        )
        .replace(
            "\\",
            "_",
        )
        .replace(
            " ",
            "_",
        )
    )

    # ============================================================
    # HTTP RESPONSE
    # ============================================================

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    return response